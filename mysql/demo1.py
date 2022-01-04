'''
Created on 2021-12-18
@author:zhangwu
Project:连接MySQL
'''

import pymysql
import openpyxl
import os


wb = openpyxl.Workbook(write_only=True)
ws = wb.create_sheet("data")
os.chdir(r"D:\pythontest\pytest\demo")


conn_drds_bic_prd = pymysql.connect(host='10.68.72.102', user='chizhou',
                                    password='EzM6v3uu', db='drds_bic_prd', charset='utf8')

conn_drds_smc_clc_prd = pymysql.connect(
    host='10.68.72.177', user='chizhou', password='EzM6v3uu', db='drds_smc_clc_prd', charset='utf8')


conn_drds_inc_prd = pymysql.connect(host='10.68.72.78', user='chizhou',
                                    password='EzM6v3uu', db='drds_inc_prd', charset='utf8')

conn_drds_smc_clc_prd = pymysql.connect(host='10.68.72.177', user='chizhou',
                                         password='EzM6v3uu', db='drds_smc_clc_prd', charset='utf8')

cursor_daiyuguanli = conn_drds_smc_clc_prd.cursor()
cursor_jibenxinxi = conn_drds_bic_prd.cursor()
cursor_caobao = conn_drds_inc_prd.cursor()
cursor_daiyuguanli = conn_drds_smc_clc_prd.cursor()

sql1 = "select table_name from information_schema.tables where table_schema='drds_bic_prd'"
sql2 = 'select * from drds_bic_prd.psn_info_b where PSN_NAME = "章小金"'


cursor_jibenxinxi.execute(sql2)

datas = cursor_jibenxinxi.fetchall()
for data in datas:
    # print(data[0])
    psn_no = data[0]

    # sql_1 = f'select * from psn_trt_info_d where psn_no = {psn_no} and cum_type_code = "Z999981"'

    sql_1 = f'select * from psn_insu_d where psn_no = {psn_no}'
    cursor_caobao.execute(sql_1)
    datas_1 = cursor_caobao.fetchall()
    for data1 in datas_1:
        # print(list(data))
        data1 = list(data1)

        data1.append(data[2])

        ws.append(data1)

wb.save("data.xlsx")
wb.close()
cursor_jibenxinxi.close()
conn_drds_bic_prd.close()
